"""
Cozeva HCC V28 validation page object.

Ports the standalone HCC V28 validation automation into the RTVS POM.
Settings come from RTVS_HCC_* env vars (see load_hcc_settings).
Login/driver/profile are provided by suite fixtures.
"""
from __future__ import annotations

import csv
import datetime
import os
import random
import time
import traceback
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from selenium.common.exceptions import (
    ElementClickInterceptedException,
    ElementNotInteractableException,
    ElementNotSelectableException,
    ElementNotVisibleException,
    NoSuchElementException,
    TimeoutException,
)
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.wait import WebDriverWait

from core.base_page import BasePage
from core.config import Config
from core.helpers import Helpers


# =============================================================================
# Measure / report constants
# =============================================================================

MEASURES = {
    400: "Review of Chronic Conditions",
    551: "Review of Suspect conditions",
    553: "HCC Score",
    554: "Review of ACA Chronic Conditions",
    555: "Review of ACA Suspect Conditions",
    556: "ACA HCC Score",
    552: "HCC Recapture",
    516: "Review of Medicaid Chronic Conditions",
    526: "Review of Medicaid Suspect Conditions",
    558: "RAF Score",
    557: "HCC Opportunity",
}

# Short aliases used in Summary "Found issues" narrative (extend when adding measures)
MEASURE_SHORT_NAMES = {
    "Review of Chronic Conditions": "RCCB",
    "Review of Chronic Conditions Version 28": "RCCB",
    "Review of Suspect conditions": "RSC",
    "HCC Score": "HCC Score",
    "Review of ACA Chronic Conditions": "ACA RCC",
    "Review of ACA Suspect Conditions": "ACA RSC",
    "ACA HCC Score": "ACA HCC Score",
    "HCC Recapture": "HCC Recapture",
    "Review of Medicaid Chronic Conditions": "Medicaid RCC",
    "Review of Medicaid Suspect Conditions": "Medicaid RSC",
    "RAF Score": "RAF Score",
    "HCC Opportunity": "HCC Opportunity",
    "HCC Efficiency": "HCC Efficiency",
}

# Fill colors reused for Summary sheet
FILL_PASS = PatternFill("solid", fgColor="BAD366")
FILL_FAIL = PatternFill("solid", fgColor="FF707A")
FILL_SKIP = PatternFill("solid", fgColor="FCD44D")
FILL_OTHER = PatternFill("solid", fgColor="D9D9D9")
FILL_HEADER = PatternFill("solid", fgColor="BDD7EE")
FILL_OK_BANNER = PatternFill("solid", fgColor="C6EFCE")
FILL_ISSUE_BANNER = PatternFill("solid", fgColor="FFC7CE")

# Measures that trigger provider CSV num/den export validation
EXPORT_VALIDATE_MEASURES = [
    "Review of Chronic Conditions",
    "Review of Chronic Conditions (Risk Adjustment Version 24)",
    "One-Year Recapture Rate",
    "One-Year Recapture Rate (Risk Adjustment Version 24)",
]

# Score-type names used when branching MSPL score math
SCORE_MEASURES = [
    "HCC Score",
    "ACA HCC Score",
    "RAF Score",
    "RAF Score (Version 28)",
    "RAF Score (ESRD Model)",
]

# ACA rate-type names used when branching MSPL gap math
ACA_RATE_MEASURES = [
    "One-Year ACA HCC Recapture Rate",
    "Review of ACA Chronic Condition",
    "One-Year ACA HCC Suspect Rate",
    "Review of ACA Suspect Condition",
]

# Common Selenium exceptions caught around fragile UI clicks
SELENIUM_EXC = (
    NoSuchElementException,
    ElementNotInteractableException,
    ElementClickInterceptedException,
    ElementNotVisibleException,
    TimeoutException,
    ElementNotSelectableException,
)


@dataclass
class HCCSettings:
    year: str = "2026"
    selected_lob: str = ""
    provider_count: int = 5
    patient_count: int = 2
    patient_dashboard: str = "No"
    patient_flag: int = 0
    selected_checklist: list[int] = field(default_factory=list)
    float_tolerance: float = 0.015
    perf_tolerance: float = 0.02
    export_validate_measures: list[str] = field(default_factory=lambda: list(EXPORT_VALIDATE_MEASURES))
    csv_download_dir: str = ""
    report_dir: str = ""
    lane_id: str | None = None


def load_hcc_settings(
    *,
    csv_download_dir: str | Path,
    report_dir: str | Path,
    lane_id: str | None = None,
) -> HCCSettings:
    """Load HCC run parameters from RTVS_HCC_* env vars with script-compatible defaults."""
    measures_raw = os.getenv("RTVS_HCC_MEASURES", "").strip()
    if measures_raw:
        selected = [int(x.strip()) for x in measures_raw.split(",") if x.strip()]
    else:
        selected = list(MEASURES.keys())

    patient_dashboard = os.getenv("RTVS_HCC_PATIENT_DASHBOARD", "No").strip() or "No"
    return HCCSettings(
        year=os.getenv("RTVS_HCC_YEAR", "2026").strip() or "2026",
        selected_lob=os.getenv("RTVS_HCC_LOB", "").strip(),
        provider_count=int(os.getenv("RTVS_HCC_PROVIDER_COUNT", "5") or 5),
        patient_count=int(os.getenv("RTVS_HCC_PATIENT_COUNT", "2") or 2),
        patient_dashboard=patient_dashboard,
        patient_flag=1 if patient_dashboard.lower() == "yes" else 0,
        selected_checklist=selected,
        float_tolerance=float(os.getenv("RTVS_HCC_FLOAT_TOLERANCE", "0.015") or 0.015),
        perf_tolerance=float(os.getenv("RTVS_HCC_PERF_TOLERANCE", "0.02") or 0.02),
        export_validate_measures=list(EXPORT_VALIDATE_MEASURES),
        csv_download_dir=str(csv_download_dir),
        report_dir=str(report_dir),
        lane_id=str(lane_id) if lane_id is not None else None,
    )


def values_match(a, b, tolerance=0.015):
    """Compare numbers with optional absolute tolerance."""
    try:
        return abs(float(a) - float(b)) <= float(tolerance)
    except (TypeError, ValueError):
        return a == b


def pass_fail(condition):
    """Return Passed/Failed string from a boolean condition."""
    return "Passed" if condition else "Failed"


def addition(total, row, switch, arr):
    """Add CSV column totals by header suffix, driven by switch code."""
    suffix_map = {
        1: ("Gaps", int),
        2: ("Conditions", int),
        3: ("Disconfirms", int),
        4: ("Clinical RAF", float),
        5: ("Potential RAF", float),
        6: ("Coded RAF", float),
        7: ("Numerator", float),
        8: ("Denominator", float),
    }
    if switch not in suffix_map:
        return total
    suffix, caster = suffix_map[switch]
    try:
        for i in range(len(arr[1])):
            if str(arr[1][i]).endswith(suffix):
                total = total + caster(arr[row][i])
        return total
    except ValueError:
        print("No number detected in " + arr[1][i] + " .For row no = " + str(row + 1))
        return total


def NumDenAddition(path):
    """Sum Numerator and Denominator columns from an export CSV."""
    with open(path, newline="") as csvfile:
        rows = list(csv.reader(csvfile, delimiter=","))
    total_num = 0
    total_den = 0
    for i in range(2, len(rows)):
        total_num = addition(total_num, i, 7, rows)
        total_den = addition(total_den, i, 8, rows)
    return total_num, total_den


def csvAddition(filepath):
    """Sum Gaps/Conditions/Disconfirms/RAF columns; return counts + patient rows."""
    with open(filepath, newline="") as csvfile:
        rows = list(csv.reader(csvfile, delimiter=","))
    Gaps = Conditions = Disconfirms = Clinical = Potential = Coded = 0
    for ind in range(2, len(rows)):
        Gaps = addition(Gaps, ind, 1, rows)
        Conditions = addition(Conditions, ind, 2, rows)
        Disconfirms = addition(Disconfirms, ind, 3, rows)
        Clinical = addition(Clinical, ind, 4, rows)
        Potential = addition(Potential, ind, 5, rows)
        Coded = addition(Coded, ind, 6, rows)
    return Gaps, Conditions, Disconfirms, Clinical, Potential, Coded, (len(rows) - 2)


def sheet_color_coder(sheet, workbook, path, filename):
    """Color-code status cells and save the workbook to path+filename."""
    rows = sheet.max_row
    cols = sheet.max_column
    for i in range(2, rows + 1):
        for j in range(3, cols + 1):
            value = sheet.cell(i, j).value
            if value is None:
                continue
            text = str(value)
            if (
                text in ("Passed", "Present and Passed")
                or "is matching" in text
                or "Successfully" in text
                or "are within" in text
                or "is present" in text
            ):
                sheet.cell(i, j).fill = PatternFill("solid", fgColor="BAD366")
            elif (
                text in ("Failed", "Present but Failed")
                or "is not matching" in text
                or "are not within" in text
                or "Issue" in text
            ):
                sheet.cell(i, j).fill = PatternFill("solid", fgColor="FF707A")
            elif text in ("Unexecuted", "Skipped") or "Warning" in text:
                sheet.cell(i, j).fill = PatternFill("solid", fgColor="FCD44D")
            elif text in ("Present but not calculated", "WAD"):
                sheet.cell(i, j).fill = PatternFill("solid", fgColor="49B99C")
            elif text == "Not present":
                sheet.cell(i, j).fill = PatternFill("solid", fgColor="FCC0BB")
    os.makedirs(path, exist_ok=True)
    path_s = str(path)
    workbook.save(path_s + filename if path_s[-1:] in ("\\", "/") else os.path.join(path_s, filename))


class CozevaHCCV28Page(BasePage):
    """HCC V28 Support/Practice/Provider MSPL and optional Patient Dashboard validation."""

    SPECIFIC_MOST = (By.CLASS_NAME, "specific_most")
    EXPORT_BULK_MSPL = (By.XPATH, "//*[@data-target='datatable_bulk_filter_0_quality_registry_list']")
    EXPORT_BULK_PROVIDER_TAB = (By.XPATH, "//*[@data-target='datatable_bulk_filter_0_metric-support-prov-ls']")
    EXPORT_ALL_TO_CSV = (By.XPATH, "//*[contains(text(),'Export all to CSV')]")

    def __init__(self, driver, settings: HCCSettings | None = None):
        super().__init__(driver)
        self.settings = settings or HCCSettings()

    def _lane_download_dir(self) -> str:
        """Lane download folder used by Chrome prefs / Helpers.fetch_downloaded_file."""
        return str(Config.RTVS_DOWNLOADS_DIR / (self.settings.lane_id or ""))

    def _export_all_to_csv(self, bulk_locator, *, wait_seconds: float = 10) -> str | None:
        """
        Click bulk menu + Export all to CSV (BasePage), then resolve the file via Helpers.
        Returns downloaded path or None.
        """
        self.click_element(bulk_locator, desc="Open bulk export menu")
        self.click_element(self.EXPORT_ALL_TO_CSV, desc="Export all to CSV")
        self.sleep_code(wait_seconds)
        self.ajax_preloader_wait("After CSV export")
        path = Helpers.fetch_downloaded_file(self.settings.lane_id)
        if path and Helpers.does_file_have_data(path):
            return path
        if path and not Helpers.does_file_have_data(path):
            print(f"[HCC] Downloaded file is empty: {path}")
            Helpers.delete_file(path)
            return None
        return None

    def validate_patient_dashboard_scores(self, patient_count, ws, wb, LOB_Name, Measure, report_path, report_filename):
        """Sample patients from MSPL and compare RAF ribbons vs patient dashboard."""
        name = self.driver.find_element(By.XPATH, "//span[@class='specific_most']").text  # Provider name
        demographic_comment = actual_raf_comment = potential_raf_comment = "Unexecuted"  # Defaults
        coded_hcc_comment = potential_hcc_comment = "Unexecuted"  # Defaults
        patient_url = "-"  # Placeholder until patient opens
        pat_count = patient_count  # How many patients to sample
        WebDriverWait(self.driver, 6000).until(  # Wait for MSPL table body
            EC.presence_of_element_located((By.XPATH, "//*[@id='quality_registry_list']/tbody"))
        )
        try:
            table = (  # Collect patient rows
                self.driver.find_element(By.ID, "quality_registry_list")
                .find_element(By.TAG_NAME, "tbody")
                .find_elements(By.TAG_NAME, "tr")
            )
            if len(table) < pat_count:  # Cap sample size to available rows
                pat_count = len(table)
                print("Not enough patients in MSPL")
            if len(table) == 1 and "No data" in table[0].text:  # Empty MSPL
                print("No patients for this provider")
                ws.append([
                    name, LOB_Name, Measure,
                    "Unexecuted", "Unexecuted", "Unexecuted", "Unexecuted", "Unexecuted",
                    "MSPL is blank for this provider", self.driver.current_url,
                ])
                return
            while pat_count != 0:  # Sample until count exhausted
                pat_count = pat_count - 1  # Decrement remaining samples
                selected_patient = table[random.randint(0, len(table) - 1)]  # Random row
                self.driver.execute_script("arguments[0].scrollIntoView();", selected_patient)  # Bring into view
                if selected_patient.find_element(By.CLASS_NAME, "risk_score_gap ").text == "NA":
                    patient_url = selected_patient.find_element(By.TAG_NAME, "a").get_attribute("href")
                    ws.append([
                        name, LOB_Name, Measure,
                        "Unexecuted", "Unexecuted", "Unexecuted", "Unexecuted", "Unexecuted",
                        "MSPL HCC value is NA so no risk score should not be present", patient_url,
                    ])
                    continue  # Try another patient
                mspl_actual_raf = float(selected_patient.find_element(By.CLASS_NAME, "actual_raf ").text)
                mspl_clinical_raf = float(selected_patient.find_element(By.CLASS_NAME, "clinical_raf ").text)
                mspl_potential_raf = float(selected_patient.find_element(By.CLASS_NAME, "potential_raf ").text)
                demographic_calculated = round(mspl_actual_raf - mspl_clinical_raf, 3)  # Demo = actual - clinical
                potential_hcc_calculated = round(mspl_potential_raf - demographic_calculated, 3)  # Pot HCC
                time.sleep(1)
                patient_link = selected_patient.find_element(By.TAG_NAME, "a")  # Patient hyperlink
                self.driver.execute_script("arguments[0].click();", patient_link)  # Open dashboard tab
                time.sleep(1)
                self.driver.switch_to.window(self.driver.window_handles[1])  # Focus new tab
                self.ajax_preloader_wait()
                WebDriverWait(self.driver, 6000).until(
                    EC.presence_of_element_located((By.CLASS_NAME, "patient_header_wrapper"))
                )
                patient_url = self.driver.current_url  # Capture dashboard URL for report
                print("Patient Dashboard URL: " + self.driver.current_url)
                score = self.driver.find_element(By.XPATH, "//*[@data-tooltip='Demographic RAF']").text
                if " -" not in score:  # Non-null demographic ribbon
                    dashboard_demographic = float(
                        self.driver.find_element(By.XPATH, "//*[@data-tooltip='Demographic RAF']").text
                    )
                    dashboard_actual_raf = float(
                        self.driver.find_element(By.XPATH, "//*[@data-tooltip='Actual RAF']").text
                    )
                    dashboard_potential_raf = float(
                        self.driver.find_element(By.XPATH, "//*[@data-tooltip='Potential RAF']").text.replace("/ ", "")
                    )
                    dashboard_coded_hcc = float(
                        self.driver.find_element(By.XPATH, "//*[@data-tooltip='Coded HCC']").text
                    )
                    dashboard_potential_hcc = float(
                        self.driver.find_element(By.XPATH, "//*[@data-tooltip='Potential HCC']").text.replace("/ ", "")
                    )
                    demographic_comment = pass_fail(values_match(demographic_calculated, dashboard_demographic, tolerance=self.settings.float_tolerance))
                    actual_raf_comment = pass_fail(values_match(mspl_actual_raf, dashboard_actual_raf, tolerance=self.settings.float_tolerance))
                    potential_raf_comment = pass_fail(values_match(mspl_potential_raf, dashboard_potential_raf, tolerance=self.settings.float_tolerance))
                    coded_hcc_comment = pass_fail(values_match(mspl_clinical_raf, dashboard_coded_hcc, tolerance=self.settings.float_tolerance))
                    potential_hcc_comment = pass_fail(values_match(potential_hcc_calculated, dashboard_potential_hcc, tolerance=self.settings.float_tolerance))
                    print("Patient Dashboard vs MSPL details")
                    print("MSPL actual raf: " + str(mspl_actual_raf))
                    print("MSPL clinical raf: " + str(mspl_clinical_raf))
                    print("MSPL potential raf: " + str(mspl_potential_raf))
                    print("Calculated demographic score: " + str(demographic_calculated))
                    print("Calculated potential HCC: " + str(potential_hcc_calculated))  # Fixed: was demographic
                    print("Dashboard actual raf: " + str(dashboard_actual_raf))
                    print("Dashboard potential raf: " + str(dashboard_potential_raf))
                    print("Dashboard coded HCC: " + str(dashboard_coded_hcc))
                    print("Dashboard potential HCC: " + str(dashboard_potential_hcc))
                    print("Dashboard demographic score: " + str(dashboard_demographic))
                    detail = (
                        "MSPL act/clin/pot="
                        + str(mspl_actual_raf) + "/" + str(mspl_clinical_raf) + "/" + str(mspl_potential_raf)
                        + " | Dash demo/act/pot/coded/potHCC="
                        + str(dashboard_demographic) + "/" + str(dashboard_actual_raf) + "/"
                        + str(dashboard_potential_raf) + "/" + str(dashboard_coded_hcc) + "/"
                        + str(dashboard_potential_hcc)
                    )
                    ws.append([
                        name, LOB_Name, Measure,
                        actual_raf_comment, potential_raf_comment, coded_hcc_comment,
                        potential_hcc_comment, demographic_comment, detail, patient_url,
                    ])
                else:  # Null HCC on dashboard
                    ws.append([
                        name, LOB_Name, Measure,
                        "Skipped", "Skipped", "Skipped", "Skipped", "Skipped",
                        "Skipped because NULL HCC value", patient_url,
                    ])
                sheet_color_coder(ws, wb, report_path, report_filename)
                self.driver.close()  # Close patient tab
                self.driver.switch_to.window(self.driver.window_handles[0])  # Back to MSPL
                time.sleep(2)
                table = (  # Refresh row list after returning
                    self.driver.find_element(By.ID, "quality_registry_list")
                    .find_element(By.TAG_NAME, "tbody")
                    .find_elements(By.TAG_NAME, "tr")
                )
        except SELENIUM_EXC as e:
            print(e)
            ws.append([
                name, LOB_Name, Measure,
                actual_raf_comment, potential_raf_comment, coded_hcc_comment,
                potential_hcc_comment, demographic_comment, "Error encountered: " + str(e), patient_url,
            ])

    def validate_provider_mspl(self, ws, wb, LOB_Name, Measure, report_path, report_filename):
        """Validate provider MSPL domain, performance, network chart, and export scores."""
        name = self.driver.find_element(By.XPATH, "//span[@class='specific_most']").text  # Provider name
        present_url = self.driver.current_url  # Restore URL after tab hops
        search_var = Measure.split(" | ")[1]  # Measure name after domain prefix
        Domain_name_MSPL = self.driver.find_element(
            By.XPATH, "//*[@class='ch metric_specific_patient_list_title']"
        ).text
        print("Domain Name MSPL page: " + Domain_name_MSPL)
        Domain_comment = pass_fail(Measure == Domain_name_MSPL)  # Domain title match
        DataToBeValidated = self.driver.find_element(By.XPATH, "//*[@class='tab']").find_elements(By.TAG_NAME, "span")
        Provider_Specific_url = self.driver.current_url
        print("Provider URL: " + Provider_Specific_url)
        DataToBeValidated_num = DataToBeValidated[0].text.replace(",", "")  # UI numerator
        print("MSPL Numerator: " + DataToBeValidated_num)
        DataToBeValidated_denum = DataToBeValidated[1].text.replace(",", "")  # UI denominator
        print("MSPL Denominator: " + DataToBeValidated_denum)
        mspl_link = self.driver.current_url  # Fallback if network tab errors
        path = self._export_all_to_csv(self.EXPORT_BULK_MSPL, wait_seconds=10)
        try:
            self.driver.find_element(By.XPATH, "//*[@class='tabs']").find_elements(By.TAG_NAME, "li")[1].click()
            self.ajax_preloader_wait()
            Performance_percentage_UI = self.driver.find_element(By.XPATH, "//*[@class='performance_value']").text.replace("%", "")
            Performance_num_UI = self.driver.find_element(By.XPATH, "//*[@class='numerator']").text.replace("Numerator: ", "")
            Performance_denum_UI = self.driver.find_element(By.XPATH, "//*[@class='denominator']").text.replace("Denominator: ", "")
            Performance_percentage_calculated = round(
                (float(Performance_num_UI) / float(Performance_denum_UI)) * 100, 4
            )
            print("Performance Tab Numerator: " + Performance_num_UI)
            print("Performance Tab Denominator: " + Performance_denum_UI)
            print("Performance Tab Percentage: " + Performance_percentage_UI)
            print("Performance calculated: " + str(Performance_percentage_calculated))
            # FIX: use abs() so under-reported UI % also fails
            Performance_comment = pass_fail(
                abs(float(Performance_percentage_UI) - float(Performance_percentage_calculated))
                < self.settings.perf_tolerance
            )
        except SELENIUM_EXC:
            Performance_comment = "The Performance tab is not clickable"
        try:
            self.driver.find_element(By.XPATH, "//*[@class='tabs']").find_elements(By.TAG_NAME, "li")[2].click()
            self.ajax_preloader_wait()
            # FIX: EC.presence_of_element_located returns a callable (always truthy); check DOM instead
            chart_present = len(self.driver.find_elements(By.XPATH, "//*[@id='network_comparison_chart']")) > 0
            Network_comment = pass_fail(chart_present)
        except SELENIUM_EXC as e:
            Network_comment = "Encountered error: " + str(e)
            traceback.print_exc()
            self.driver.get(mspl_link)
            self.ajax_preloader_wait()
        if path is None:
            ws.append([
                name, LOB_Name, Measure, Domain_comment, Performance_comment, Network_comment,
                "Failed", "No CSV found in " + self._lane_download_dir(), Provider_Specific_url,
            ])
            self.driver.get(present_url)
            return
        result = csvAddition(path)  # Aggregate export columns
        print("Total Gap Count: " + str(result[0]))
        print("Total Condition Count: " + str(result[1]))
        print("Total Disconfirm Count: " + str(result[2]))
        print("Total Clinical RAF Score: " + str(result[3]))
        print("Total Potential Score: " + str(result[4]))
        print("Total Coded Score: " + str(result[5]))
        print("Total Patient Count: " + str(result[6]))
        Helpers.delete_file(path)
        tol = self.settings.float_tolerance
        if search_var in SCORE_MEASURES:  # RAF / HCC score measures
            DataToBeValidated_num = round(float(DataToBeValidated_denum) - float(DataToBeValidated_num), 3)
            DataToBeValidated_denum = round(float(DataToBeValidated_denum), 3)
            num = round(float(result[3] / result[6]), 3)  # Avg clinical from export
            temp = float((result[5] - result[3]) / result[6])
            denum = round(float(result[4] / result[6]) - temp, 3)
            comments = (
                "MSPL (UI): " + str(DataToBeValidated_num) + "/" + str(DataToBeValidated_denum)
                + " || MSPL (Export): " + str(num) + "/" + str(denum)
            )
            status = pass_fail(
                abs(float(DataToBeValidated_num) - num) < tol
                and abs(float(DataToBeValidated_denum) - denum) < tol
            )
        elif search_var in ACA_RATE_MEASURES:  # ACA rate / review measures
            num = result[0]
            denum = result[0] + result[1] + result[2]
            comments = (
                "MSPL (UI): " + str(DataToBeValidated_num) + "/" + str(DataToBeValidated_denum)
                + " || MSPL (Export): " + str(num) + "/" + str(denum)
            )
            status = pass_fail(
                int(DataToBeValidated_num) == int(DataToBeValidated_denum) - int(result[1]) - int(result[2])
            )
        else:  # Default gap/condition style measures
            num = result[0]
            denum = result[0] + result[1] + result[2]
            comments = (
                "MSPL (UI): " + str(DataToBeValidated_num) + "/" + str(DataToBeValidated_denum)
                + " || MSPL (Export): " + str(num) + "/" + str(denum)
            )
            status = pass_fail(
                float(DataToBeValidated_num) == float(result[0])
                and float(DataToBeValidated_denum) == float(result[0] + result[1] + result[2])
            )
        ws.append([
            name, LOB_Name, Measure, Domain_comment, Performance_comment, Network_comment,
            status, comments, Provider_Specific_url,
        ])
        sheet_color_coder(ws, wb, report_path, report_filename)
        self.driver.get(present_url)  # Restore provider page

    def collect_measure_urls(self, LOB, checklist, URL_name_list, num_den_list, customer):
        """Collect measure URLs and names for the selected measure IDs on registry."""
        for i in checklist:  # Each selected measure id
            try:
                Measure_Specific_url = self.driver.find_element(By.XPATH, "//*[@id=" + str(i) + "]//a").get_attribute("href")
                Measure = self.driver.find_element(By.XPATH, "//*[@id=" + str(i) + "]//*[@class='met-name top']").text
                Domain_name_registry = self.driver.find_element(
                    By.XPATH, "//*[@id=" + str(i) + "]/..//*[@class='group-name-wrapper']"
                ).text
                Measure = Domain_name_registry + " | " + Measure  # Domain | Measure
                URL_name_list.append(Measure_Specific_url)  # URL
                URL_name_list.append(Measure)  # Display name
                URL_name_list.append(LOB)  # LOB for this entry
            except Exception:
                print("No measure of id " + str(i))
        return URL_name_list, num_den_list

    def navigate_support_tabs(self, Tabs, ws, wb, measure, LOB, num_den_list, provider_count, id_list, report_path, report_filename):
        """Walk Practices/Providers/Patients/Performance tabs; collect practice & provider links."""
        name = self.driver.find_element(By.XPATH, "//span[@class='specific_most']").text  # Context name
        Tabs = self.driver.find_element(By.XPATH, "//*[@id='qt-mt-support-ls']").find_elements(By.TAG_NAME, "li")
        practice_link = ""  # First practice drill-down URL
        provider_link = []  # Sampled provider URLs
        pro_count = provider_count  # How many providers to sample
        for i in range(0, len(Tabs) - 1):  # Skip last (usually unused) tab
            comments = ""
            try:
                Tabs[i].click()
                self.ajax_preloader_wait()
                WebDriverWait(self.driver, 300).until(EC.presence_of_element_located((By.CLASS_NAME, "tab")))
                field = self.driver.find_element(By.XPATH, "//*[@class='handler active']").text  # Active tab label
                print(field)
                if field == "Practices":
                    try:
                        ListRow = (
                            self.driver.find_element(By.XPATH, "//*[@id='metric-support-prac-ls']")
                            .find_element(By.TAG_NAME, "tbody")
                            .find_elements(By.TAG_NAME, "tr")
                        )
                        if "No data available" in ListRow[0].text and len(ListRow) == 1:
                            comments = "No data available in practice tab"
                        else:
                            Row = ListRow[0]  # Use first practice row
                            practice_link = Row.find_elements(By.TAG_NAME, "a")[1].get_attribute("href")
                            comments = self.driver.find_element(By.XPATH, "//*[@class='dataTables_info']").text
                    except NoSuchElementException:
                        comments = "Tab faced an error while opening"
                if field == "Providers":
                    try:
                        ListRow = (
                            self.driver.find_element(By.XPATH, "//*[@id='metric-support-prov-ls']")
                            .find_element(By.TAG_NAME, "tbody")
                            .find_elements(By.TAG_NAME, "tr")
                        )
                        if "No data available" in ListRow[0].text and len(ListRow) == 1:
                            comments = "No data available in provider tab"
                        else:
                            if len(ListRow) < pro_count:
                                pro_count = len(ListRow)
                                print("Has less number of providers than specified")
                            while pro_count > 0:
                                Row = ListRow[random.randint(0, (len(ListRow) - pro_count))]
                                link = Row.find_elements(By.TAG_NAME, "a")[1].get_attribute("href")
                                provider_link.append(link)
                                pro_count -= 1
                            comments = self.driver.find_element(By.XPATH, "//*[@class='dataTables_info']").text
                    except NoSuchElementException:
                        comments = "Tab faced an error while opening"
                    try:
                        if measure in id_list:  # Optional export num/den vs registry
                            path = self._export_all_to_csv(self.EXPORT_BULK_PROVIDER_TAB, wait_seconds=2)
                            if path:
                                ws.append([
                                    name, LOB, measure, "Successfully downloaded export file", "-",
                                    self.driver.current_url,
                                ])
                                total_num, total_den = NumDenAddition(path)
                                total_num = int(total_num)
                                total_den = int(total_den)
                                print(total_num)
                                print(total_den)
                                for j in range(0, len(num_den_list) - 1, 4):
                                    if LOB in num_den_list[j + 2] and measure in num_den_list[j + 3]:
                                        if total_num == num_den_list[j] and total_den == num_den_list[j + 1]:
                                            ws.append([
                                                name, LOB, measure,
                                                "The registry num/denum count is matching with export",
                                                "Registry:" + str(num_den_list[j]) + "/" + str(num_den_list[j + 1])
                                                + "|| Export:" + str(total_num) + "/" + str(total_den),
                                                self.driver.current_url,
                                            ])
                                        else:
                                            ws.append([
                                                name, LOB, measure,
                                                "The registry num/denum count is not matching with export",
                                                "Registry:" + str(num_den_list[j]) + "/" + str(num_den_list[j + 1])
                                                + "|| Export:" + str(total_num) + "/" + str(total_den),
                                                self.driver.current_url,
                                            ])
                                Helpers.delete_file(path)
                            else:
                                ws.append([
                                    name, LOB, measure,
                                    "Failed to download export file",
                                    "No CSV found in " + self._lane_download_dir(),
                                    self.driver.current_url,
                                ])
                    except NoSuchElementException:
                        comments = "Tab faced an error while opening"
                if field == "Patients":
                    try:
                        ListRow = (
                            self.driver.find_element(By.XPATH, "//*[@id='metric-support-pat-ls']")
                            .find_element(By.TAG_NAME, "tbody")
                            .find_elements(By.TAG_NAME, "tr")
                        )
                        if "No data available" in ListRow[0].text and len(ListRow) == 1:
                            comments = "No data available in patients tab"
                        else:
                            comments = self.driver.find_element(By.XPATH, "//*[@class='dataTables_info']").text
                    except NoSuchElementException:
                        comments = "Tab faced an error while opening"
                if field == "Performance Statistics":
                    try:
                        Performance_percentage_UI = self.driver.find_element(
                            By.XPATH, "//*[@class='performance_value']"
                        ).text.replace("%", "")
                        Performance_num_UI = self.driver.find_element(
                            By.XPATH, "//*[@class='numerator']"
                        ).text.replace("Numerator: ", "")
                        Performance_denum_UI = self.driver.find_element(
                            By.XPATH, "//*[@class='denominator']"
                        ).text.replace("Denominator: ", "")
                        try:
                            Performance_percentage_calculated = round(
                                (float(Performance_num_UI) / float(Performance_denum_UI)) * 100, 4
                            )
                            # FIX: abs() so both over/under deltas fail correctly
                            if abs(float(Performance_percentage_UI) - float(Performance_percentage_calculated)) < self.settings.perf_tolerance:
                                comments = "The performance percentage matches with num/denum value"
                            else:
                                comments = "Performance percentage does not match with num/denum value"
                        except Exception as e:
                            print(e)
                            comments = "Denominator value is zero"
                    except Exception:
                        comments = "Tab faced an error while opening"
                ws.append([name, LOB, measure, "Successfully navigated to " + field, comments, self.driver.current_url])
                Tabs = self.driver.find_element(By.ID, "qt-mt-support-ls").find_elements(By.TAG_NAME, "li")  # Re-query after DOM refresh
            except SELENIUM_EXC:
                ws.append([
                    name, LOB, measure,
                    "Issue: Tab encountered error for tab no.: " + str(i + 1), "NA", self.driver.current_url,
                ])
        sheet_color_coder(ws, wb, report_path, report_filename)
        return practice_link, provider_link

    def measure_short_name(self, measure_text):
        """Map full Domain | Measure text to a short alias for Summary findings."""
        raw = str(measure_text or "").strip()
        # Prefer the part after " | " when present
        leaf = raw.split(" | ")[-1].strip() if " | " in raw else raw
        if leaf in MEASURE_SHORT_NAMES:
            return MEASURE_SHORT_NAMES[leaf]
        for key, short in MEASURE_SHORT_NAMES.items():  # Partial match fallback
            if key.lower() in leaf.lower() or leaf.lower() in key.lower():
                return short
        return leaf or "Unknown measure"

    def create_report_workbook(self, customer_value, year):
        """Create workbook with Summary first, then Support / Practice / Provider / Patient sheets."""
        style_name = "hcc_v28_header"
        try:
            header = NamedStyle(name=style_name)
            header.font = Font(bold=True)
            header.border = Border(bottom=Side(border_style="thin"))
            header.alignment = Alignment(horizontal="center", vertical="center")
        except ValueError:
            # NamedStyle already registered in this process
            header = NamedStyle(name=style_name + "_" + str(id(self)))
            header.font = Font(bold=True)
            header.border = Border(bottom=Side(border_style="thin"))
            header.alignment = Alignment(horizontal="center", vertical="center")

        wb = Workbook()
        # Summary is created first so it opens as the lead sheet
        ws = wb.active
        ws.title = "Summary"
        ws.append(["Sheet", "Passed", "Failed", "Skipped/Unexecuted", "Other", "Total rows"])
        for cell in ws[1]:
            cell.style = header
            cell.fill = FILL_HEADER

        sheet_name = (customer_value + " Support Registry " + str(year))[:31]  # Excel 31-char limit
        wb.create_sheet(sheet_name)
        ws = wb[sheet_name]
        ws.append(["Customer Name", "LOB", "Measure", "Status", "Comments", "URL"])
        for cell in ws[1]:
            cell.style = header

        wb.create_sheet("Practice Registry")
        ws = wb["Practice Registry"]
        ws.append(["Practice Name", "LOB", "Measure", "Status", "Comments", "URL"])
        for cell in ws[1]:
            cell.style = header

        wb.create_sheet("Provider MSPL")
        ws = wb["Provider MSPL"]
        ws.append([
            "Provider Name", "LOB", "Measure", "Domain Name Check",
            "Performance Statistics Check", "Network Comparison Check",
            "Risk Score Check", "Comments", "Provider URL",
        ])
        for cell in ws[1]:
            cell.style = header

        wb.create_sheet("Patient Dashboard")
        ws = wb["Patient Dashboard"]
        ws.append([
            "Provider name", "LOB", "Measure", "Actual RAF Check", "Potential RAF Check",
            "Coded HCC Check", "Potential HCC Check", "Demographic Score Check", "Comment", "URL",
        ])
        for cell in ws[1]:
            cell.style = header

        # Keep Summary as the first sheet index
        wb.move_sheet("Summary", offset=-len(wb.sheetnames) + 1)
        return wb, sheet_name, header

    def _row_status_bucket(self, joined):
        """Classify a joined row string into fail / skip / pass / other."""
        keywords_pass = ("Passed", "matching", "Successfully", "are within", "is present", "successful")
        keywords_fail = ("Failed", "not matching", "are not within", "Issue", "mismatch", "Error encountered")
        keywords_skip = ("Unexecuted", "Skipped", "Warning")
        if any(k in joined for k in keywords_fail):
            return "fail"
        if any(k in joined for k in keywords_skip):
            return "skip"
        if any(k in joined for k in keywords_pass):
            return "pass"
        return "other"

    def _collect_summary_issues(self, wb):
        """
        Build numbered finding lines from detail sheets.
        Groups by issue type + short measure name, listing providers/patients.
        """
        findings = []  # Ordered list of narrative strings

        # --- Provider MSPL: risk score / export mismatches ---
        mspl_by_measure = {}  # short_measure -> set(provider names)
        domain_by_measure = {}
        perf_by_measure = {}
        network_by_measure = {}
        if "Provider MSPL" in wb.sheetnames:
            sheet = wb["Provider MSPL"]
            for r in range(2, sheet.max_row + 1):
                provider = str(sheet.cell(r, 1).value or "").strip() or "Unknown provider"
                measure = str(sheet.cell(r, 3).value or "")
                short = self.measure_short_name(measure)
                domain = str(sheet.cell(r, 4).value or "")
                perf = str(sheet.cell(r, 5).value or "")
                network = str(sheet.cell(r, 6).value or "")
                risk = str(sheet.cell(r, 7).value or "")
                comments = str(sheet.cell(r, 8).value or "")
                if risk == "Failed" or "not matching" in comments.lower():
                    mspl_by_measure.setdefault(short, set()).add(provider)
                if domain == "Failed":
                    domain_by_measure.setdefault(short, set()).add(provider)
                if perf == "Failed":
                    perf_by_measure.setdefault(short, set()).add(provider)
                if network == "Failed" or network.startswith("Encountered error"):
                    network_by_measure.setdefault(short, set()).add(provider)

        for short, providers in sorted(mspl_by_measure.items()):
            findings.append(
                "MSPL vs UI export count mismatch for measure "
                + short
                + " for following providers: "
                + ", ".join(sorted(providers))
                + "."
            )
        for short, providers in sorted(domain_by_measure.items()):
            findings.append(
                "Domain name mismatch on Provider MSPL for measure "
                + short
                + " for following providers: "
                + ", ".join(sorted(providers))
                + "."
            )
        for short, providers in sorted(perf_by_measure.items()):
            findings.append(
                "Performance Statistics mismatch on Provider MSPL for measure "
                + short
                + " for following providers: "
                + ", ".join(sorted(providers))
                + "."
            )
        for short, providers in sorted(network_by_measure.items()):
            findings.append(
                "Network Comparison check failed on Provider MSPL for measure "
                + short
                + " for following providers: "
                + ", ".join(sorted(providers))
                + "."
            )

        # --- Patient Dashboard: any RAF check Failed ---
        patients_by_measure = {}  # short -> set("provider | url")
        if "Patient Dashboard" in wb.sheetnames:
            sheet = wb["Patient Dashboard"]
            for r in range(2, sheet.max_row + 1):
                provider = str(sheet.cell(r, 1).value or "").strip() or "Unknown provider"
                measure = str(sheet.cell(r, 3).value or "")
                short = self.measure_short_name(measure)
                checks = [
                    str(sheet.cell(r, c).value or "") for c in range(4, 9)
                ]  # Actual..Demographic
                url = str(sheet.cell(r, 10).value or "").strip()
                if any(c == "Failed" for c in checks):
                    label = provider if not url or url == "-" else provider + " (" + url + ")"
                    patients_by_measure.setdefault(short, set()).add(label)

        for short, patients in sorted(patients_by_measure.items()):
            findings.append(
                "Patient dashboard RAF score mismatch observed for measure "
                + short
                + " for following patients: "
                + ", ".join(sorted(patients))
                + "."
            )

        # --- Support / Practice registry narrative issues ---
        for name in wb.sheetnames:
            if name in ("Summary", "Provider MSPL", "Patient Dashboard"):
                continue
            sheet = wb[name]
            by_measure = {}
            for r in range(2, sheet.max_row + 1):
                entity = str(sheet.cell(r, 1).value or "").strip() or "Unknown"
                measure = str(sheet.cell(r, 3).value or "")
                short = self.measure_short_name(measure)
                status = str(sheet.cell(r, 4).value or "")
                comments = str(sheet.cell(r, 5).value or "")
                joined = status + " | " + comments
                if "not matching" in joined.lower() or status.startswith("Issue") or "Failed" in status:
                    by_measure.setdefault(short, set()).add(entity)
            for short, entities in sorted(by_measure.items()):
                label = "Support Registry" if "Support Registry" in name else name
                findings.append(
                    "Issue observed in "
                    + label
                    + " for measure "
                    + short
                    + " affecting: "
                    + ", ".join(sorted(entities))
                    + "."
                )

        return findings

    def update_summary_sheet(self, wb):
        """Rebuild Summary (first sheet): color-coded counts + Found issues / no issues banner."""
        if "Summary" not in wb.sheetnames:
            return
        # Always place Summary first for validation review
        try:
            wb.move_sheet("Summary", offset=-wb.sheetnames.index("Summary"))
        except Exception:
            pass

        ws_sum = wb["Summary"]
        # Clear entire sheet then rewrite
        ws_sum.delete_rows(1, ws_sum.max_row)

        # --- Section 1: count table ---
        headers = ["Sheet", "Passed", "Failed", "Skipped/Unexecuted", "Other", "Total rows"]
        ws_sum.append(headers)
        for cell in ws_sum[1]:
            cell.font = Font(bold=True)
            cell.fill = FILL_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center")

        total_failed = 0
        for name in wb.sheetnames:
            if name == "Summary":
                continue
            sheet = wb[name]
            passed = failed = skipped = other = 0
            for r in range(2, sheet.max_row + 1):
                row_vals = [str(c.value) if c.value is not None else "" for c in sheet[r]]
                bucket = self._row_status_bucket(" | ".join(row_vals))
                if bucket == "fail":
                    failed += 1
                elif bucket == "skip":
                    skipped += 1
                elif bucket == "pass":
                    passed += 1
                else:
                    other += 1
            total = max(0, sheet.max_row - 1)
            total_failed += failed
            ws_sum.append([name, passed, failed, skipped, other, total])
            row_idx = ws_sum.max_row
            # Color code count cells: green / red / yellow / gray
            ws_sum.cell(row_idx, 2).fill = FILL_PASS if passed else PatternFill()
            ws_sum.cell(row_idx, 3).fill = FILL_FAIL if failed else FILL_PASS
            ws_sum.cell(row_idx, 4).fill = FILL_SKIP if skipped else PatternFill()
            ws_sum.cell(row_idx, 5).fill = FILL_OTHER if other else PatternFill()

        # Widen columns for readability (A holds findings text too)
        ws_sum.column_dimensions["A"].width = 100
        for col in ("B", "C", "D", "E", "F"):
            ws_sum.column_dimensions[col].width = 18

        # --- Section 2: overall verdict + numbered findings ---
        ws_sum.append([])  # blank spacer row
        findings = self._collect_summary_issues(wb)

        if not findings and total_failed == 0:
            ws_sum.append(["Overall result", "No issues to report"])
            banner_row = ws_sum.max_row
            ws_sum.cell(banner_row, 1).font = Font(bold=True, color="006100")
            ws_sum.cell(banner_row, 2).font = Font(bold=True, color="006100")
            ws_sum.cell(banner_row, 1).fill = FILL_OK_BANNER
            ws_sum.cell(banner_row, 2).fill = FILL_OK_BANNER
        else:
            ws_sum.append(["Overall result", "Issues observed — see Found issues below"])
            banner_row = ws_sum.max_row
            ws_sum.cell(banner_row, 1).font = Font(bold=True, color="9C0006")
            ws_sum.cell(banner_row, 2).font = Font(bold=True, color="9C0006")
            ws_sum.cell(banner_row, 1).fill = FILL_ISSUE_BANNER
            ws_sum.cell(banner_row, 2).fill = FILL_ISSUE_BANNER

            ws_sum.append([])
            ws_sum.append(["Found issues:"])
            header_row = ws_sum.max_row
            ws_sum.cell(header_row, 1).font = Font(bold=True, size=12)
            ws_sum.cell(header_row, 1).fill = FILL_ISSUE_BANNER

            if not findings:
                # Count table showed fails but collector found none — generic line
                findings = [
                    "Failures were detected in sheet counts; review Failed cells in detail sheets."
                ]

            for idx, text in enumerate(findings, start=1):
                ws_sum.append([str(idx) + ". " + text])
                issue_row = ws_sum.max_row
                ws_sum.cell(issue_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
                ws_sum.cell(issue_row, 1).fill = FILL_FAIL
                ws_sum.row_dimensions[issue_row].height = max(30, 15 * (1 + len(text) // 100))




    def run_hcc_v28_validation(self, *, customer_value: str, config_assists=None) -> Path:
        """
        Run HCC V28 validation from an already-logged-in Support Registries context.
        Returns path to the generated Excel report.
        """
        settings = self.settings
        os.makedirs(settings.csv_download_dir, exist_ok=True)
        os.makedirs(settings.report_dir, exist_ok=True)

        self.ajax_preloader_wait("Post-login settle before HCC validation")
        customer_name = self.driver.find_element(By.CLASS_NAME, "specific_most").text

        wb, sheet_name, header = self.create_report_workbook(str(customer_value), settings.year)
        path1 = settings.report_dir
        if not path1.endswith(("\\", "/")):
            path1 = path1 + os.sep
        formatted_date = datetime.datetime.now().strftime("%m-%d %H-%M")
        safe_customer = "".join(c if c.isalnum() or c in (" ", "-", "_") else "_" for c in customer_name).strip()
        filename = f"{safe_customer}_HCC Check MY{settings.year}_{formatted_date}.xlsx"

        if config_assists:
            config_assists.add_log_update(
                message=(
                    f"HCC V28 start | year={settings.year} lob={settings.selected_lob or 'ALL'} "
                    f"measures={settings.selected_checklist} providers={settings.provider_count} "
                    f"patients={settings.patient_count} dashboard={settings.patient_dashboard}"
                ),
                driver=self.driver,
                status="IN_PROGRESS",
            )

        if config_assists:
            config_assists.add_log_heartbeat(
                f"Applying {settings.year} year filter and LOB selection "
                f"({settings.selected_lob or 'ALL LOBs'})",
                driver=self.driver,
                status="STARTED",
            )

        self.driver.find_element(By.XPATH, "//*[@id='qt-filter-label']").click()
        Quarter_list = self.driver.find_element(By.XPATH, "//*[@id='filter-quarter']").find_elements(By.TAG_NAME, "li")
        for quarter in Quarter_list:
            if quarter.text == settings.year:
                quarter.click()
                break

        LOB_list = self.driver.find_element(By.XPATH, "//*[@id='filter-lob']").find_elements(By.TAG_NAME, "li")
        URL_name_registry_list = []
        num_den_value_list = []
        num_den_prac_list = []
        selected_LOB = settings.selected_lob
        flag = 0

        for i in range(0, len(LOB_list)):
            flag += 1
            LOB_Name = LOB_list[i].text
            print("LOB Name for " + settings.year + ": " + LOB_Name)
            try:
                if selected_LOB != "":
                    if LOB_Name == selected_LOB:
                        LOB_list[i].click()
                    else:
                        continue
                else:
                    LOB_list[i].click()
            except ElementNotInteractableException:
                continue
            self.driver.find_element(By.ID, "reg-filter-apply").click()
            self.ajax_preloader_wait("After LOB apply: " + LOB_Name)
            if self.driver.find_element(By.XPATH, "//*[@id='conti_enroll']").is_selected():
                self.driver.find_element(By.XPATH, "//*[@class='cont_disc_toggle']").click()
            print("Current support registry URL: " + self.driver.current_url)
            if config_assists:
                config_assists.add_log_heartbeat(
                    f"Collecting selected measures on Support Registry for LOB: {LOB_Name}",
                    driver=self.driver,
                    status="IN_PROGRESS",
                )
            URL_name_registry_list, num_den_value_list = self.collect_measure_urls(
                LOB_Name,
                settings.selected_checklist,
                URL_name_registry_list,
                num_den_value_list,
                customer_name,
            )
            while flag < len(LOB_list):
                self.driver.find_element(By.XPATH, "//*[@id='qt-filter-label']").click()
                LOB_list = self.driver.find_element(By.XPATH, "//*[@id='filter-lob']").find_elements(By.TAG_NAME, "li")
                break
            print(URL_name_registry_list)
            print(num_den_value_list)

        id_list = settings.export_validate_measures
        provider_count = settings.provider_count
        patient_count = settings.patient_count
        patient_flag = settings.patient_flag

        total_measures = len(URL_name_registry_list) // 3
        if config_assists:
            config_assists.add_log_heartbeat(
                f"Collected {total_measures} measure(s) across LOBs; starting MSPL validation",
                driver=self.driver,
                status="STARTED",
            )

        for i in range(0, len(URL_name_registry_list), 3):
            url = URL_name_registry_list[i]
            measure = URL_name_registry_list[i + 1]
            LOB = URL_name_registry_list[i + 2]
            measure_index = (i // 3) + 1
            print(url)
            if config_assists:
                config_assists.add_log_heartbeat(
                    f"Validating measure {measure_index}/{total_measures}: {measure} [{LOB}]",
                    driver=self.driver,
                    status="IN_PROGRESS",
                )
            self.driver.get(url)
            self.ajax_preloader_wait("Open measure MSPL: " + measure)
            time.sleep(3)
            ws = wb[sheet_name]
            MSPL_name = self.driver.find_element(
                By.XPATH, "//*[@class='ch metric_specific_patient_list_title valign-wrapper']"
            ).text
            print(MSPL_name)
            if URL_name_registry_list[i + 1] in MSPL_name:
                status = "Domain name successful validation"
                comment = "Checked for Registry vs MSPL"
                case_status = "PASSED"
            else:
                status = "Domain name mismatch"
                comment = "Not matching for Registry vs MSPL"
                case_status = "FAILED"
            ws.append([
                self.driver.find_element(By.XPATH, "//span[@class='specific_most']").text,
                LOB,
                measure,
                status,
                comment,
                url,
            ])
            if config_assists:
                config_assists.add_log_test_case(
                    message=f"Support Registry domain check: {measure} / {LOB}",
                    test_case_id="HCC_V28_01",
                    status=case_status,
                    driver=self.driver,
                    comment=comment,
                )
            WebDriverWait(self.driver, 300).until(EC.presence_of_element_located((By.ID, "qt-mt-support-ls")))
            WebDriverWait(self.driver, 300).until(EC.presence_of_element_located((By.CLASS_NAME, "tab")))
            Tabs = self.driver.find_element(By.ID, "qt-mt-support-ls").find_elements(By.TAG_NAME, "li")
            practice_link, provider_link_supports = self.navigate_support_tabs(
                Tabs,
                ws,
                wb,
                measure,
                LOB,
                num_den_value_list,
                provider_count,
                id_list,
                path1,
                filename,
            )
            sheet_color_coder(ws, wb, path1, filename)
            ws = wb["Practice Registry"]
            try:
                if config_assists:
                    config_assists.add_log_heartbeat(
                        f"Opening Practice Registry for {measure} [{LOB}]",
                        driver=self.driver,
                        status="IN_PROGRESS",
                    )
                self.driver.get(practice_link)
                self.ajax_preloader_wait("Open practice registry")
                WebDriverWait(self.driver, 300).until(
                    EC.presence_of_element_located((By.XPATH, "//*[@data-target='qt-reg-nav-filters']"))
                )
                search_var = measure.split(" | ")[1]
                print(search_var)
                time.sleep(2)
                self.driver.find_element(
                    By.XPATH,
                    "//*[contains(@class, 'met-name top') and contains(text(), '"
                    + search_var
                    + "')] /.. /.. /.. /..",
                ).click()
                self.ajax_preloader_wait("Practice measure click")
                Tabs = self.driver.find_element(By.ID, "qt-mt-support-ls").find_elements(By.TAG_NAME, "li")
                self.navigate_support_tabs(
                    Tabs,
                    ws,
                    wb,
                    measure,
                    LOB,
                    num_den_prac_list,
                    provider_count,
                    id_list,
                    path1,
                    filename,
                )
                sheet_color_coder(ws, wb, path1, filename)
                print("Provider links")
                print(provider_link_supports)
                total_providers = len(provider_link_supports)
                for provider_idx, provider_link_support in enumerate(provider_link_supports, start=1):
                    self.driver.get(provider_link_support)
                    self.ajax_preloader_wait("Open provider MSPL")
                    ws = wb["Provider MSPL"]
                    if config_assists:
                        config_assists.add_log_heartbeat(
                            f"Validating Provider MSPL {provider_idx}/{total_providers} "
                            f"for {measure} [{LOB}]",
                            driver=self.driver,
                            status="IN_PROGRESS",
                        )
                    self.validate_provider_mspl(ws, wb, LOB, measure, path1, filename)
                    sheet_color_coder(ws, wb, path1, filename)
                    self.ajax_preloader_wait("After provider MSPL check")
                    ws = wb["Patient Dashboard"]
                    print("Patient flag: " + str(patient_flag))
                    if patient_flag != 0:
                        if config_assists:
                            config_assists.add_log_heartbeat(
                                f"Validating Patient Dashboard RAF scores "
                                f"(sampling {patient_count} patient(s)) for {measure} [{LOB}]",
                                driver=self.driver,
                                status="IN_PROGRESS",
                            )
                        self.validate_patient_dashboard_scores(
                            patient_count, ws, wb, LOB, measure, path1, filename
                        )
                    else:
                        print("Patient Dashboard Score check skipped")
                    sheet_color_coder(ws, wb, path1, filename)
            except Exception as e:
                print(e)
                traceback.print_exc()
                ws.append([customer_name, LOB, measure, "Issue: Blank MSPL encountered", "NA", self.driver.current_url])
                sheet_color_coder(ws, wb, path1, filename)
                if config_assists:
                    config_assists.add_log_test_case(
                        message=f"Practice/Provider walk failed: {measure} / {LOB}",
                        test_case_id="HCC_V28_02",
                        status="FAILED",
                        driver=self.driver,
                        comment=str(e),
                    )

        if config_assists:
            config_assists.add_log_heartbeat(
                "Generating summary sheet and finalizing HCC V28 report",
                driver=self.driver,
                status="IN_PROGRESS",
            )
        self.update_summary_sheet(wb)
        sheet_color_coder(wb[sheet_name], wb, path1, filename)
        if path1.endswith(("\\", "/")):
            report_path = Path(path1 + filename)
        else:
            report_path = Path(path1) / filename
        print("Report saved to: " + str(report_path))
        if config_assists:
            config_assists.add_log_update(
                message=f"HCC V28 Excel report saved: {report_path}",
                driver=self.driver,
                status="Success",
            )
        return report_path
